"""Comprehensive Report Generation Services (PDF, Excel & CSV)."""
import io
import csv
from decimal import Decimal
from datetime import datetime
from django.utils import timezone
from apps.billing.models import Invoice
from apps.inventory.models import Ingredient
from apps.orders.models import Order, OrderStatus
from apps.core.utils import format_inr

# OpenPyXL for Excel
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ReportLab for PDF
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

def generate_sales_pdf(start_date, end_date, branch=None):
    """
    Generates a statutory Sales Audit PDF report using ReportLab.
    Returns a BytesIO buffer.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    elements = []
    styles = getSampleStyleSheet()

    # Header title
    title_style = ParagraphStyle(
        'ReportTitle',
        parent=styles['Heading1'],
        fontSize=18,
        textColor=colors.HexColor('#d97706'),
        alignment=1,  # Centered
        spaceAfter=10
    )
    elements.append(Paragraph("DineFlow Enterprise ERP — Sales Audit Report", title_style))
    
    sub_text = f"Period: {start_date} to {end_date}"
    if branch:
        sub_text += f" | Branch: {branch.name}"
    elements.append(Paragraph(sub_text, styles['Normal']))
    elements.append(Spacer(1, 15))

    # Table data
    qs = Order.objects.filter(
        created_at__date__gte=start_date,
        created_at__date__lte=end_date,
        status=OrderStatus.COMPLETED
    )
    if branch:
        qs = qs.filter(branch=branch)

    data = [['Order #', 'Date/Time', 'Table', 'Subtotal', 'Tax', 'Grand Total']]
    tot_sub = Decimal('0.00')
    tot_tax = Decimal('0.00')
    tot_grand = Decimal('0.00')

    for o in qs.order_by('created_at'):
        data.append([
            o.order_number,
            o.created_at.strftime('%d/%m %H:%M'),
            o.table.table_number if o.table else 'Counter',
            f"{o.subtotal:.2f}",
            f"{o.tax_amount:.2f}",
            f"{o.grand_total:.2f}"
        ])
        tot_sub += o.subtotal
        tot_tax += o.tax_amount
        tot_grand += o.grand_total

    data.append(['TOTALS', '', '', f"{tot_sub:.2f}", f"{tot_tax:.2f}", f"{tot_grand:.2f}"])

    t = Table(data, colWidths=[110, 85, 75, 80, 75, 95])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#090b10')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.HexColor('#f59e0b')),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('ALIGN', (3, 1), (-1, -1), 'RIGHT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#f1f5f9')),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
    ]))
    elements.append(t)

    doc.build(elements)
    buffer.seek(0)
    return buffer

def generate_gstr1_excel(start_date, end_date, branch=None):
    """
    Generates Indian GST GSTR-1 Sales Return in Excel (.xlsx).
    Returns a BytesIO buffer.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "GSTR-1 B2C Summary"

    # Header styling
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="F59E0B")
    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    headers = [
        "Invoice #", "Invoice Date", "GSTIN", "Place of Supply",
        "Taxable Value (INR)", "CGST (2.5%)", "SGST (2.5%)", "IGST (5.0%)",
        "Total Tax (INR)", "Invoice Total (INR)"
    ]
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    qs = Invoice.objects.filter(
        created_at__date__gte=start_date,
        created_at__date__lte=end_date,
        is_paid=True
    )
    if branch:
        qs = qs.filter(branch=branch)

    row_num = 2
    for inv in qs.order_by('created_at'):
        ws.append([
            inv.invoice_number,
            inv.created_at.strftime('%Y-%m-%d'),
            inv.restaurant_gstin,
            inv.branch.state if inv.branch else "Telangana",
            float(inv.taxable_subtotal),
            float(inv.cgst_amount),
            float(inv.sgst_amount),
            float(inv.igst_amount),
            float(inv.total_tax_amount),
            float(inv.grand_total)
        ])
        for col_idx in range(5, 11):
            ws.cell(row=row_num, column=col_idx).alignment = align_right
        row_num += 1

    # Auto column width
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

def generate_inventory_valuation_excel(branch=None):
    """
    Generates Inventory Stock Valuation in Excel (.xlsx).
    Returns a BytesIO buffer.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stock Valuation"

    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="10B981")
    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    headers = [
        "SKU Code", "Ingredient Name", "Category", "Current Stock", "Unit",
        "Reorder Level", "Unit Cost (INR)", "Total Value (INR)", "Stock Status"
    ]
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    qs = Ingredient.objects.select_related('category', 'branch').all()
    if branch:
        qs = qs.filter(branch=branch)

    row_num = 2
    for ing in qs.order_by('category__name', 'name'):
        status_str = "CRITICAL LOW" if ing.is_low_stock else "NORMAL"
        ws.append([
            ing.code,
            ing.name,
            ing.category.name if ing.category else "General",
            float(ing.current_stock),
            ing.unit,
            float(ing.minimum_stock_level),
            float(ing.unit_cost),
            float(ing.stock_value),
            status_str
        ])
        for col_idx in [4, 6, 7, 8]:
            ws.cell(row=row_num, column=col_idx).alignment = align_right
        row_num += 1

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

def generate_sales_csv(start_date, end_date, branch=None):
    """Generates CSV string of sales orders between dates."""
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        'Order Number', 'Order Date', 'Branch', 'Order Type', 'Table',
        'Subtotal (INR)', 'Discount (INR)', 'GST Tax (INR)', 'Grand Total (INR)', 'Status'
    ])

    qs = Order.objects.filter(
        created_at__date__gte=start_date,
        created_at__date__lte=end_date
    )
    if branch:
        qs = qs.filter(branch=branch)

    for o in qs.order_by('created_at'):
        writer.writerow([
            o.order_number,
            o.created_at.strftime('%Y-%m-%d %H:%M'),
            o.branch.name if o.branch else '',
            o.get_order_type_display(),
            o.table.table_number if o.table else 'Counter',
            f"{o.subtotal:.2f}",
            f"{o.discount_amount:.2f}",
            f"{o.tax_amount:.2f}",
            f"{o.grand_total:.2f}",
            o.status
        ])

    return output.getvalue()

def generate_inventory_valuation_csv(branch=None):
    """Generates CSV string of raw material inventory valuation and stock levels."""
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        'SKU Code', 'Ingredient Name', 'Category', 'Current Stock', 'Unit',
        'Reorder Level', 'Unit Cost (INR)', 'Total Stock Value (INR)', 'Is Low Stock'
    ])

    qs = Ingredient.objects.select_related('category', 'branch').all()
    if branch:
        qs = qs.filter(branch=branch)

    for ing in qs.order_by('category__name', 'name'):
        writer.writerow([
            ing.code,
            ing.name,
            ing.category.name if ing.category else 'Uncategorized',
            f"{ing.current_stock:.3f}",
            ing.unit,
            f"{ing.minimum_stock_level:.3f}",
            f"{ing.unit_cost:.2f}",
            f"{ing.stock_value:.2f}",
            'YES' if ing.is_low_stock else 'NO'
        ])

    return output.getvalue()

def generate_gstr1_b2c_csv(month, year, branch=None):
    """Generates Indian GST statutory GSTR-1 B2C register report in CSV."""
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        'Invoice Number', 'Invoice Date', 'Branch GSTIN', 'Place of Supply',
        'Taxable Value (INR)', 'CGST (2.5%)', 'SGST (2.5%)', 'IGST (5.0%)',
        'Total Tax (INR)', 'Invoice Total (INR)'
    ])

    qs = Invoice.objects.filter(
        created_at__year=year,
        created_at__month=month,
        is_paid=True
    )
    if branch:
        qs = qs.filter(branch=branch)

    for inv in qs.order_by('created_at'):
        writer.writerow([
            inv.invoice_number,
            inv.created_at.strftime('%Y-%m-%d'),
            inv.restaurant_gstin,
            inv.branch.state if inv.branch else 'Telangana',
            f"{inv.taxable_subtotal:.2f}",
            f"{inv.cgst_amount:.2f}",
            f"{inv.sgst_amount:.2f}",
            f"{inv.igst_amount:.2f}",
            f"{inv.total_tax_amount:.2f}",
            f"{inv.grand_total:.2f}"
        ])

    return output.getvalue()
